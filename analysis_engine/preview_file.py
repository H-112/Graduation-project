"""
文件预览工具 — 读取CSV/Excel的前5行
用法: python3 preview_file.py <filepath>
"""
import sys
import json
import csv
from pathlib import Path

def main():
    filepath = Path(sys.argv[1])
    ext = filepath.suffix.lower()

    try:
        if ext == '.csv':
            with open(filepath, 'r', encoding='utf-8-sig') as f:
                reader = csv.reader(f)
                headers = next(reader, [])
                rows = []
                for row in reader:
                    rows.append([str(c)[:80] for c in row])
                    if len(rows) >= 5:
                        # Count remaining rows
                        remaining = sum(1 for _ in reader)
                        break
                else:
                    remaining = 0
                total = len(rows)

            print(json.dumps({
                "headers": headers,
                "rows": rows[:5],
                "totalRows": total + remaining,
            }, ensure_ascii=False))

        elif ext in ('.xlsx', '.xls'):
            import openpyxl
            wb = openpyxl.load_workbook(filepath, read_only=True)
            ws = wb.active
            headers = [str(ws.cell(1, c).value or '') for c in range(1, ws.max_column + 1)]
            rows = []
            for r in range(2, min(ws.max_row + 1, 7)):
                row = [str(ws.cell(r, c).value or '')[:80] for c in range(1, ws.max_column + 1)]
                rows.append(row)
            wb.close()

            print(json.dumps({
                "headers": headers,
                "rows": rows,
                "totalRows": ws.max_row - 1,
            }, ensure_ascii=False))

    except Exception as e:
        print(json.dumps({"error": str(e), "headers": [], "rows": [], "totalRows": 0}))

if __name__ == '__main__':
    main()
