"""
LLM 问卷结构分析器 — 用 DeepSeek 自动识别题型、人口学字段、量表分组
解决了纯规则引擎的语义盲区（如把院系/书院误判为文本题）
用法: python3 llm_structure_analyzer.py <input_file>
输出: JSON 结构映射到 stdout
"""
import json
import csv
import sys
import os
from pathlib import Path
from collections import Counter
from openai import OpenAI

def progress(msg: str) -> None:
    print(f"[PROGRESS] {msg}", flush=True)

# ── 环境变量 ──────────────────────────────────
def _load_env():
    env_file = Path(__file__).parent.parent / ".env"
    if env_file.exists():
        with open(env_file, 'r') as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith('#') and '=' in line:
                    k, v = line.split('=', 1)
                    k, v = k.strip(), v.strip().strip('"').strip("'")
                    if k not in os.environ:
                        os.environ[k] = v

_load_env()
DEEPSEEK_API_KEY = os.environ.get("DEEPSEEK_API_KEY", "")
DEEPSEEK_BASE_URL = "https://api.deepseek.com"

def get_client():
    if not DEEPSEEK_API_KEY:
        sys.exit("错误: 未设置 DEEPSEEK_API_KEY")
    return OpenAI(api_key=DEEPSEEK_API_KEY, base_url=DEEPSEEK_BASE_URL)

# ── 加载与采样 ───────────────────────────────

def load_file(filepath: Path):
    ext = filepath.suffix.lower()
    if ext == '.csv':
        with open(filepath, 'r', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            return reader.fieldnames or [], list(reader)
    elif ext in ('.xlsx', '.xls'):
        import openpyxl
        wb = openpyxl.load_workbook(filepath)
        ws = wb.active
        headers = [ws.cell(1, c).value or f"Col{c}" for c in range(1, ws.max_column + 1)]
        rows = []
        for r in range(2, ws.max_row + 1):
            row = {}
            for c, h in enumerate(headers, 1):
                val = ws.cell(r, c).value
                row[h] = str(val).strip() if val is not None else ""
            rows.append(row)
        return headers, rows
    else:
        sys.exit(f"不支持的文件格式: {ext}")

def sample_values(values: list[str], n=5):
    """采样非空值，优先不同值"""
    non_empty = [v.strip() for v in values if v.strip()]
    if len(non_empty) <= n:
        return non_empty
    # 取前n个不同的值
    seen = set()
    samples = []
    for v in non_empty:
        if v not in seen:
            seen.add(v)
            samples.append(v)
        if len(samples) >= n:
            break
    return samples

# ── 提示词构建 ───────────────────────────────

def build_prompt(headers: list[str], col_data: dict) -> str:
    """构建问卷结构分析提示词"""
    lines = [
        "你是一位专业的问卷分析系统。请分析以下问卷的列结构，对每一列进行精确分类。",
        "",
        "## 分类规则",
        "",
        "为每一列标注以下类型之一：",
        "- **demographic**: 人口学/身份信息（性别、年级、专业、学院、书院、系、身份、住宿地、年级、年龄 等）",
        "- **single_choice**: 单选题（选项互斥，只选一个）",
        "- **multi_select**: 多选题（可选多个）",
        "- **likert**: Likert 量表题（1-5分评分，满意度/同意度等程度评价）",
        "- **text**: 文本/开放题（自由填写的长文字回答）",
        "- **skip**: 元数据/非题目（编号、序号、时间戳、IP、答题时长、提交时间、设备信息 等）",
        "",
        "## 重要判断原则",
        "1. 如果列名包含「书院」「院系」「学院」「专业」「班级」「高校」「大学」「学校」「年级」「年」「月」「年龄」「性别」且值是具体名称或个人信息（非A/B/C/D选项），属于 **demographic**，即使唯一值很多（如200+个不同大学名称）也仍然是人口学信息",
        "2. 如果值像 1-5 的数字或「非常满意/满意/一般/不满意/非常不满意」等程度词，属于 **likert**",
        "3. 如果值很简短（≤20字）且唯一值≤20个，属于 **single_choice**",
        "4. 如果值包含多选分隔符（┋ ; ；）或列名含「多选」「可多选」「多选题」等，属于 **multi_select**",
        "5. 如果值是长文本（平均>30字），属于 **text**",
        "6. 如果列名含「编号」「ID」「时间」「IP」「答题时长」「提交」等且非调查问题，属于 **skip**",
        "7. 注意：高校/大学/学校名称即使有上百个不同的值，也属于 **demographic**（人口学），不是 text。同理，年份/月份/出生年月等日期信息也属于 **demographic**",
        "",
        "## 字段信息",
        "每列提供：列名 + 采样值(最多5个) + 唯一值个数",
        "",
    ]

    for i, h in enumerate(headers, 1):
        vals = col_data.get(h, [])
        non_empty = [v for v in vals if v.strip()]
        unique_count = len(set(non_empty))
        samples = sample_values(vals, 5)

        lines.append(f"### 列{i}: {h}")
        lines.append(f"唯一值数: {unique_count}")
        lines.append(f"总非空数: {len(non_empty)}")
        if samples:
            lines.append(f"采样值:")
            for s in samples:
                lines.append(f"  - \"{s[:80]}\"")
        lines.append("")

    lines.append("---")
    lines.append("请以 JSON 格式返回分类结果。对于 type=likert 的列，还需提供 value_map（文本标签→1-5分值）。")
    lines.append("```json")
    lines.append("{")
    lines.append('  "columns": [')
    lines.append('    {"column": "列名", "type": "类型", "reason": "简要原因(≤20字)",')
    lines.append('     "value_map": {"非常满意": 5, "比较满意": 4, "一般": 3, "比较不满意": 2, "非常不满意": 1}}')
    lines.append('    ...')
    lines.append("  ]")
    lines.append("}")
    lines.append("```")
    lines.append("注意: value_map 仅 type=likert 时需要，格式为 {去前缀后的值: 1-5分值}。取 5 个最积极→最消极的层级。")

    return '\n'.join(lines)


# ── 主流程 ───────────────────────────────────

def main():
    if len(sys.argv) < 2:
        print("Usage: python3 llm_structure_analyzer.py <input_file>")
        sys.exit(1)

    infile = Path(sys.argv[1])
    progress(f"加载文件: {infile.name}")
    headers, rows = load_file(infile)

    # 取每列值（全量用于统计，LLM只用采样）
    col_values = {}
    for h in headers:
        col_values[h] = [row.get(h, '') for row in rows]

    # 只发送非空列给 LLM（节省 token）
    active_columns = []
    for h in headers:
        non_empty = sum(1 for v in col_values[h] if v.strip())
        if non_empty > 0:
            active_columns.append(h)

    progress(f"共 {len(headers)} 列，{len(active_columns)} 列有数据，发送给 LLM 分析...")

    # 分批：每批最多 40 列，防止 token 超限
    BATCH_SIZE = 40
    batches = [active_columns[i:i + BATCH_SIZE] for i in range(0, len(active_columns), BATCH_SIZE)]

    client = get_client()
    all_results = []

    for batch_idx, batch_headers in enumerate(batches):
        batch_data = {h: col_values[h] for h in batch_headers}
        prompt = build_prompt(batch_headers, batch_data)

        progress(f"分析批次 {batch_idx + 1}/{len(batches)} ({len(batch_headers)} 列)...")

        try:
            response = client.chat.completions.create(
                model="deepseek-chat",
                messages=[
                    {
                        "role": "system",
                        "content": "你是专业的问卷结构分析系统。你只返回 JSON，不返回其他内容。你的分析准确率高，能区分人口学问题、量表题、选择题、文本题和元数据列。"
                    },
                    {"role": "user", "content": prompt}
                ],
                temperature=0.1,  # 低温，提高一致性
                max_tokens=4000,
            )

            content = response.choices[0].message.content
            # 提取 JSON 块
            json_str = content
            if "```json" in json_str:
                json_str = json_str.split("```json")[1].split("```")[0]
            elif "```" in json_str:
                json_str = json_str.split("```")[1].split("```")[0]

            batch_results = json.loads(json_str)
            columns = batch_results.get('columns', [])
            all_results.extend(columns)
            progress(f"  ✓ 批次 {batch_idx + 1}: 分类 {len(columns)} 列")

        except json.JSONDecodeError as e:
            progress(f"  ✗ JSON 解析失败 (batch {batch_idx + 1}): {e}")
            progress(f"  原始回复前 500 字: {content[:500]}")
        except Exception as e:
            progress(f"  ✗ API 调用失败 (batch {batch_idx + 1}): {e}")

    # 按原有顺序输出该列，未分类的保留为 'unknown'
    result_map = {r['column']: r for r in all_results}
    output_columns = []
    for h in headers:
        if h in result_map:
            output_columns.append(result_map[h])
        else:
            non_empty = sum(1 for v in col_values.get(h, []) if v.strip())
            if non_empty > 0:
                output_columns.append({"column": h, "type": "unknown", "reason": "LLM未返回"})
            else:
                output_columns.append({"column": h, "type": "skip", "reason": "空列"})

    # 统计分类
    type_counts = Counter(c['type'] for c in output_columns)
    progress(f"分类统计: {dict(type_counts)}")

    result = {
        "file": infile.name,
        "total_columns": len(headers),
        "classification": output_columns,
        "summary": dict(type_counts),
    }

    print(json.dumps(result, ensure_ascii=False, indent=2))

if __name__ == '__main__':
    main()
