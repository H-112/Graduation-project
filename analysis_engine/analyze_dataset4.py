"""
数据集 #4 分析引擎：生成式人工智能对大学生学习方式的影响（915份）
模式1：描述性统计 + 频数分析 + 文本分析
"""
import openpyxl
import json
import math
import os
import sys
from collections import Counter
from pathlib import Path

# 从dataset5脚本导入工具函数
sys.path.insert(0, str(Path(__file__).parent))
from analyze_dataset5 import (
    load_data, describe_categorical, describe_likert,
    describe_multi_select, tokenize_chinese, STOPWORDS
)

DATA_DIR = Path(__file__).parent.parent / "data"
OUTPUT_DIR = Path(__file__).parent / "output"
FNAME = "265987486_按文本_调查：生成式人工智能对大学生学习方式的影响_915_915.xlsx"

def text_analysis(records, col_name):
    """开放题文本分析"""
    texts = []
    for r in records:
        val = r.get(col_name)
        if val is not None and str(val).strip():
            texts.append(str(val).strip())

    if not texts:
        return {"column": col_name, "total_answers": 0, "note": "No text responses"}

    all_words = []
    answer_lengths = []
    for text in texts:
        words = tokenize_chinese(text)
        words = [w for w in words if len(w) >= 2 and w not in STOPWORDS]
        all_words.extend(words)
        answer_lengths.append(len(text))

    word_counter = Counter(all_words)
    top_keywords = word_counter.most_common(50)

    avg_len = sum(answer_lengths) / len(answer_lengths)

    return {
        "column": col_name,
        "total_answers": len(texts),
        "avg_answer_length": round(avg_len, 1),
        "total_words_extracted": len(all_words),
        "unique_words": len(word_counter),
        "top_keywords": [
            {"word": w, "count": c, "rank": i+1}
            for i, (w, c) in enumerate(top_keywords)
        ]
    }

def main():
    print("=" * 60)
    print("数据集 #4 分析引擎")
    print("生成式人工智能对大学生学习方式的影响")
    print("=" * 60)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # 1. 加载数据
    print("\n[1/5] 加载数据...")
    wb = openpyxl.load_workbook(DATA_DIR / FNAME)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    records = []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        rec = {}
        for i, h in enumerate(headers):
            rec[h] = row[i] if i < len(row) else None
        records.append(rec)

    print(f"  总记录数: {len(records)}")
    print(f"  总字段数: {len(headers)}")

    # 2. 人口学分析
    print("\n[2/5] 人口学分析...")
    demographics = {}
    demographics['gender'] = describe_categorical(records, '1.您的性别：')
    demographics['grade'] = describe_categorical(records, '2.您的年级：')
    demographics['major'] = describe_categorical(records, '3.您的专业')
    demographics['university_type'] = describe_categorical(records, '5.您所在的高校类型是:')
    demographics['family_location'] = describe_categorical(records, '7.您的家庭所在地位于：')

    # 论文和科研参与
    demographics['published_papers'] = describe_categorical(records, '(1)6.您发表或在写论文的数量___')
    demographics['research_projects'] = describe_categorical(records, '(2)，参与科研项目的数量___')

    # 3. GenAI 使用行为
    print("\n[3/5] GenAI使用行为分析...")
    genai = {}
    genai['start_time'] = describe_categorical(records, '1.您是从什么时候开始使用GenAI?')
    genai['tools_used'] = describe_categorical(records, '2.您使用过哪些GenAI?（多选）')
    genai['familiarity'] = describe_categorical(records, '3.您对GenAI的熟悉程度是:')
    genai['proficiency'] = describe_categorical(records, '4.我能够熟练运用GenAI产出我需要的答案:')
    genai['frequency'] = describe_categorical(records, '5.您使用GenAI的频率:')
    genai['duration_per_use'] = describe_categorical(records, '6.您每次使用GenAI的时间:')

    # 使用场景
    scene_headers = [
        '7.您在以下学习场景中使用GenAI的频率:—课堂学习场景',
        '课后作业场景',
        '科研项目场景',
        '毕业论文场景',
    ]
    existing_scenes = [h for h in scene_headers if h in headers]
    genai['scene_usage'] = {}
    for sh in existing_scenes:
        genai['scene_usage'][sh] = describe_categorical(records, sh)

    # 学校限制
    genai['school_policy'] = describe_categorical(records, '8.学校/学院是否对使用生成式人工智能会有限制条款:')

    # 4. Likert 量表分析
    print("\n[4/5] Likert量表分析...")
    likert = {}

    # 课堂学习场景使用GenAI
    classroom_items = [
        '①课堂学习场景：—我会使用GenAI提前预习课程内容。',
        '我使用GenAI回答课堂上老师的提问。',
        '我会使用GAI查询课堂上并未详述的信息，以扩充知识范围。',
        '我使用GenAI生成课堂讨论的想法和观点。',
    ]
    existing_classroom = [h for h in classroom_items if h in headers]
    likert['classroom_genai_usage'] = describe_likert(records, existing_classroom, '课堂GenAI使用')

    # 课后作业
    homework_items = [
        '②课后作业场景：—我会使用GenAI辅助完成课程作业，生成相关资料。',
        '我会使用GenAI帮助规划作业的整体结构和思路。',
        '我会直接将GenAI的产生内容作为作业或项目的一部分上交。',
        '我会让GenAI评价作业并给予反馈。',
    ]
    existing_homework = [h for h in homework_items if h in headers]
    likert['homework_genai_usage'] = describe_likert(records, existing_homework, '课后作业GenAI使用')

    # 科研项目
    research_items = [
        '③科研项目场景：—我会使用GenAI协助研究项目的选题。',
        '我会使用GenAI查找文献并阅读文献。',
        '我会使用GenAI优化研究项目的整体结构和实验设计。',
        '我会使用GenAI分析科研数据，帮助得出结论或解释结果。',
    ]
    existing_research = [h for h in research_items if h in headers]
    likert['research_genai_usage'] = describe_likert(records, existing_research, '科研GenAI使用')

    # 毕业论文
    thesis_items = [
        '④毕业论文场景：—我会使用GenAI查找并阅读毕业论文的文献。',
        '我会使用GenAI查找并整理数据。',
        '我会使用GenAI为毕业论文的章节设计和内容提供建议。',
        '我会使用GenAI润色毕业论文的语言（如语法，词语等）。',
        '我会使用GenAI调整毕业论文的格式。',
        '我会使用GenAI生成毕业论文的初步草稿并进一步修改。',
    ]
    existing_thesis = [h for h in thesis_items if h in headers]
    likert['thesis_genai_usage'] = describe_likert(records, existing_thesis, '毕业论文GenAI使用')

    # 学习方法量表（Q2 18道题）
    learning_method_items = [
        '2.请如实、快速作答，以便准确描述自己的实际学习方法。—我经常难以理解我需要记住的事物。',
        '当我阅读文章或书籍时，我会努力自己理解作者的真正意图。',
        '我会认真安排我的学习时间，以充分利用它。',
        '我发现有大部分工作都不怎么有趣或相关。',
        '我会均匀地分配整个学期的学习任务，而不是等到最后一刻。',
        '在解决问题或作业前，我会首先尝试理解背后的深层原因。',
        '我能够在需要时迅速开始学习工作。',
        '我学习的很多内容似乎毫无意义，就像是零散的碎片。',
        '我投入大量努力学习，因为我决心要做好。',
        '当我学习新主题时，我会尝试在脑海中整合所有的观点。',
        '我并不觉得激励自己学习是困难的。',
        '我经常会对在讲课中听到的或在书籍中读到的内容提出疑问。',
        '我认为自己在考试复习时非常有系统和组织。',
        '我经常觉得自己快要被我们必须应对的大量材料淹没了。',
        '课本或文章中的观点经常会引发我长时间的思考。',
        '我不太确定讲课中什么是重要的，所以我尽可能多地记笔记。',
        '阅读时，我会仔细检查细节，看它们如何与所说的内容相符。',
        '我经常担心自己是否能够妥善应对学习工作。',
    ]
    existing_lm = [h for h in learning_method_items if h in headers]
    likert['learning_methods'] = describe_likert(records, existing_lm, '学习方法')

    # 课程/教学偏好
    course_pref_items = [
        '3.请如实、快速作答，以便准确描述自己偏好的课程和教学。—明确告诉我们应该记下什么的讲师',
        '鼓励我们独立思考并展示他们自己的思维方式的讲师',
        '让我可以展示我自己对课程材料的思考的考试',
        '只需利用我们讲座笔记中提供的材料的考试',
        '非常明确指示我们应该阅读哪些书籍的课程',
        '鼓励我们自己大量阅读相关主题的课程',
        '挑战你并提供超越讲座内容的解释的书籍',
        '提供明确事实和信息并且容易学习的书籍',
    ]
    existing_cp = [h for h in course_pref_items if h in headers]
    likert['course_preferences'] = describe_likert(records, existing_cp, '课程偏好')

    # GenAI对学习的影响
    genai_impact_items = [
        '五、生成式人工智能（GenAI）带来的学习方式改变请如实、快速作答，以便准确描述GenAI对自己学习的影响。—使用GenAI提高了我的成绩，与使用前相比有显著提升。',
        '我发现经常使用GenAI的同学成绩更好。',
        '我对GenAI在帮助解决学习问题方面的效果感到满意。',
        'GenAI在合理解决学习难题方面表现出很高的准确性和实用性。',
        '使用GenAI提升了我的自主学习能力，使我在课堂外能够更有效地掌握新知识。',
        '我发现GenAI为我提供了启发性想法，提高了我的创新能力。',
        'GenAI拓宽了我看待问题的角度，增强了我的批判性思维能力。',
        '使用GenAI之后我的思考问题能力下降。',
        '使用GenAI之后我的解决问题能力下降。',
        '使用GenAI让我依赖其来完成日常学习、项目或考试。',
        '我认为生成式人工智能不应该在学习中被使用。',
    ]
    existing_gi = [h for h in genai_impact_items if h in headers]
    likert['genai_impact'] = describe_likert(records, existing_gi, 'GenAI对学习的影响')

    # 5. 开放题文本分析
    print("\n[5/5] 开放题文本分析...")
    text = {}

    q1 = text_analysis(records, '1.您认为生成式人工智能对大学生学习方式有哪些影响?')
    text['genai_impact_opinion'] = q1

    q2 = text_analysis(records, '2.对于使用生成式人工智能改革学习方式，您有什么建议?')
    text['genai_reform_suggestions'] = q2

    # 组装输出
    print("\n生成结果文件...")
    result = {
        "dataset": "生成式人工智能对大学生学习方式的影响",
        "file": FNAME,
        "total_records": len(records),
        "total_fields": len(headers),
        "analysis_mode": "quick_overview",
        "demographics": demographics,
        "genai_usage": genai,
        "likert_scales": likert,
        "text_analysis": text,
    }

    output_path = OUTPUT_DIR / "dataset4_analysis.json"
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(f"  ✓ 完整结果: {output_path} ({output_path.stat().st_size / 1024:.1f} KB)")

    # 生成摘要
    summary = generate_summary(result)
    summary_path = OUTPUT_DIR / "dataset4_summary.md"
    with open(summary_path, 'w', encoding='utf-8') as f:
        f.write(summary)
    print(f"  ✓ 摘要报告: {summary_path}")

    print("\n" + "=" * 60)
    print("分析完成!")
    print("=" * 60)
    return result

def generate_summary(result):
    lines = []
    lines.append("# 数据集 #4 分析摘要：生成式人工智能对大学生学习方式的影响")
    lines.append(f"\n**样本量**: {result['total_records']} 份")
    lines.append(f"**字段数**: {result['total_fields']}")

    lines.append("\n## 1. 样本构成")
    for key in ['1.您的性别：', '2.您的年级：', '5.您所在的高校类型是:']:
        if key in result['demographics']:
            d = result['demographics'][key]
            lines.append(f"\n### {d['column']}")
            for item in d['distribution'][:6]:
                lines.append(f"- {item['label']}: {item['count']} ({item['percentage']}%)")

    lines.append("\n## 2. GenAI 使用行为")
    for key in ['1.您是从什么时候开始使用GenAI?', '5.您使用GenAI的频率:', '3.您对GenAI的熟悉程度是:']:
        if key in result['genai_usage']:
            d = result['genai_usage'][key]
            lines.append(f"\n### {d['column']}")
            for item in d['distribution'][:5]:
                lines.append(f"- {item['label']}: {item['count']} ({item['percentage']}%)")

    lines.append("\n## 3. 核心量表得分")
    for group_name, items in result['likert_scales'].items():
        if items:
            lines.append(f"\n### {items[0]['group'] if items else group_name}")
            for item in items[:6]:
                short_name = item['column'].split('—')[-1] if '—' in item['column'] else item['column'].split(':')[-1]
                lines.append(f"- {short_name.strip()}: 均值={item['mean']}, SD={item['std']}")

    lines.append("\n## 4. 开放题关键词")
    for key, ta in result['text_analysis'].items():
        lines.append(f"\n### {ta.get('column', key)}")
        lines.append(f"  有效回答: {ta.get('total_answers', 0)} 份")
        if 'top_keywords' in ta:
            top10 = ta['top_keywords'][:10]
            lines.append(f"  Top 10: {', '.join(w['word'] for w in top10)}")

    return '\n'.join(lines)

if __name__ == '__main__':
    main()
